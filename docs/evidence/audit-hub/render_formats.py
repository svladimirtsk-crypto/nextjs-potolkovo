from pathlib import Path
import re,json,base64,html
import mistune
from bs4 import BeautifulSoup,NavigableString
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path('/home/user/audit')
md=(B/'POTOLKOVO-FINAL-AUDIT-TZ-V6.md').read_text()
markdown=mistune.create_markdown(plugins=['table','strikethrough','url'])
body=markdown(md)
soup=BeautifulSoup(body,'html.parser')
for tag in list(soup.find_all(['h1','h2','h3']))[:3]:tag.decompose()
headings=[]
for i,h in enumerate(soup.find_all('h2')):
 h['id']='section-'+str(i);headings.append((h['id'],h.get_text()))
# Cover and contents use real metadata, no invented endorsement.
cover='''<section class="cover"><div class="eyebrow">ПОТОЛКОВО / INDEPENDENT REVIEW</div><div class="version">AUDIT + SPEC / V6</div><h1>От расчёта потолка<br>к проверенному<br>умному освещению</h1><p class="sub">Консолидированный аудит и техническое задание</p><div class="rule"></div><p>Клиентские пути · EKS · комплектация · цена<br>Разработка · дизайн · маркетинг · эксплуатация</p><div class="cover-bottom"><b>14 октября 2026 · Москва</b><br>Проверенный код: 7f1b5cf · TZv5 / PT-020<br><span>Предложение к утверждению. Не отчёт о внедрении.</span></div></section>'''
toc='<section class="toc"><div class="eyebrow">НАВИГАЦИЯ</div><h2>Содержание</h2><p>Для владельца: разделы 0, 3, 13, 15. Для исполнителей: 7–14.</p><ol>'+''.join(f'<li><a href="#{id}">{html.escape(t)}</a></li>' for id,t in headings)+'</ol><div class="note"><b>Важная граница:</b> выполнен аудит исходников, локальной production-сборки и доступных данных EKS. Продакшен и физические smart-комплекты не сертифицированы этой проверкой.</div></section>'
figs=[('360-home.png','A15 · Мобильная шапка','Viewport 360 px; document.scrollWidth 383 px. Источник: overflow360.json. Нужно исправить компоновку header, а не маскировать overflow у body.'),('quiz-home-5.png','A08 · Готовый расчёт и idle-подсказка','После подтверждения стандартных значений есть итог 18 000 ₽, но остаётся приглашение ответить на вопросы. Подтверждение должно считаться ответом.'),('led-c02.png','A01 · Контроллеры размечены метрами','C02 и соседние товары управления показываются как /м. Отдельным серверным тестом подтверждено принятие 0,5 контроллера за 1004 ₽.'),('390-home.png','Контекст · Главная на 390 px','Контрольный первый экран. На этой ширине общего переполнения не обнаружено; это не опровергает отдельный дефект на 360 px.')]
gallery='<section class="visual"><h2>Визуальные доказательства</h2>'
for name,title,caption in figs:
 raw=base64.b64encode((B/'evidence/screens'/name).read_bytes()).decode()
 gallery+=f'<figure><h3>{title}</h3><img src="data:image/png;base64,{raw}"><figcaption>{caption}<br><code>evidence/screens/{name}</code></figcaption></figure>'
gallery+='</section>'
css='''
@page {size:A4; margin:19mm 17mm 19mm;}
*{box-sizing:border-box}html{scroll-behavior:smooth}body{font-family:Arial,"DejaVu Sans",sans-serif;color:#1e293b;font-size:10pt;line-height:1.5;margin:0;background:#fff}h1,h2,h3{color:#142c43;line-height:1.2}h1{font-size:31pt}h2{font-size:21pt;margin:0 0 20px;padding-top:10px}h3{font-size:13pt;margin:24px 0 9px}h4{font-size:11pt}p{margin:8px 0}a{color:#176f83;text-decoration:none;overflow-wrap:anywhere}code{font-family:"DejaVu Sans Mono",monospace;font-size:.87em;background:#eff3f5;padding:1px 3px;overflow-wrap:anywhere}pre{padding:16px;background:#eff3f5;white-space:pre-wrap;overflow-wrap:anywhere;font-size:8.4pt;line-height:1.5}pre code{padding:0}blockquote{margin:16px 0;border-left:4px solid #008e9a;background:#ecf6f7;padding:12px 18px}ul,ol{padding-left:23px}li{margin:5px 0}table{border-collapse:collapse;width:100%;margin:13px 0 20px;table-layout:auto;font-size:8.6pt;line-height:1.4}th{background:#142c43;color:white;text-align:left;font-weight:600}th,td{padding:8px;border-bottom:1px solid #d6e1e6;vertical-align:top;overflow-wrap:anywhere}tr:nth-child(even) td{background:#f2f6f8}thead{display:table-header-group}tr{break-inside:avoid}hr{height:1px;border:0;background:#d3e0e7;margin:25px 0}.cover{height:251mm;position:relative;padding:13mm 4mm;color:#142c43}.eyebrow{letter-spacing:2px;font-size:9pt;font-weight:700;color:#008593}.version{font-size:12pt;margin-top:39mm;font-weight:bold;letter-spacing:3px;color:#008593}.cover h1{font-size:35pt;line-height:1.17;margin:14px 0 18px}.sub{font-size:16pt;max-width:130mm}.rule{height:5px;width:60px;background:#00a3aa;margin:23px 0}.cover-bottom{position:absolute;bottom:11mm;line-height:1.7;font-size:11pt}.cover-bottom span{font-size:9pt;color:#637889}.toc{padding-top:15mm;break-before:page;break-after:page}.toc h2{margin-top:10mm}.toc ol{list-style:none;padding-left:0}.toc li{margin:7px 0;font-size:11pt}.note{margin-top:20px;background:#edf6f7;padding:16px;border-top:2px solid #008c98}article h2{margin-top:26px}article #section-3,article #section-7,article #section-11,article #section-12,article #section-14{break-before:page}h2,h3,h4{break-after:avoid}figure{break-inside:avoid;margin:18px 0 30px;display:block}figure img{width:auto;max-width:100%;height:140mm;object-fit:contain;object-position:left center}figcaption{font-size:9pt;color:#526577;max-width:150mm;margin-top:8px}.visual{break-before:page}.visual figure:nth-of-type(3){break-before:page}
@media screen{body{max-width:980px;margin:auto;padding:25px 40px}.cover{height:850px}.cover-bottom{bottom:35px}article h2{padding-top:40px;margin-top:50px;border-top:1px solid #d3e0e7}.toc{margin:30px 0 70px}.visual figure img{max-height:650px;height:auto}.cover h1{font-size:46px}table{display:block;overflow-x:auto}}
@media print{.cover{break-after:page}.toc{break-before:auto}p,li{orphans:2;widows:2}}
'''
html_doc='<!doctype html><html lang="ru"><meta charset="utf-8"><title>ПОТОЛКОВО — аудит и ТЗ v6</title><style>'+css+'</style><body>'+cover+toc+'<article>'+str(soup)+'</article>'+gallery+'</body></html>'
(B/'POTOLKOVO-FINAL-AUDIT-TZ-V6.html').write_text(html_doc)
# Word document
D=Document();sec=D.sections[0];sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
styles=D.styles
styles['Normal'].font.name='Calibri';styles['Normal'].font.size=Pt(10)
styles['Normal'].paragraph_format.space_after=Pt(5)
for n,size in [('Title',34),('Heading 1',22),('Heading 2',14),('Heading 3',12)]:
 styles[n].font.name='Calibri';styles[n].font.size=Pt(size);styles[n].font.color.rgb=RGBColor.from_string('142C43')
 styles[n].paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.page_break_before=True
h=sec.header.paragraphs[0];h.text='ПОТОЛКОВО  /  КОНСОЛИДИРОВАННЫЙ АУДИТ И ТЗ v6';h.style='Caption';h.runs[0].font.color.rgb=RGBColor.from_string('6E8390')
foot=sec.footer.paragraphs[0];foot.text='14.10.2026 · commit 7f1b5cf                            '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
D.add_paragraph('ПОТОЛКОВО / AUDIT + SPEC V6','Subtitle');D.add_paragraph('\n\n')
D.add_paragraph('От расчёта потолка\nк проверенному\nумному освещению','Title')
D.add_paragraph('Консолидированный аудит и техническое задание','Subtitle')
D.add_paragraph('Клиентские пути · EKS · комплектация · цена\nРазработка · дизайн · маркетинг · эксплуатация')
D.add_paragraph('\n\n14 октября 2026 · Москва\nПроверенный код: 7f1b5cf · TZv5 / PT-020\nПредложение к утверждению. Не отчёт о внедрении.')
D.add_heading('Содержание',1)
for _,title in headings:D.add_paragraph(title)
D.add_paragraph('Навигация Word: Вид → Область навигации. Полные критерии исполнения также представлены в Excel.','Caption')
def hyperlink(p,label,url):
 part=p.part;rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'176F83');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);el.append(r);p._p.append(el)
def inline(p,node,bold=False,italic=False):
 if isinstance(node,NavigableString):
  r=p.add_run(str(node));r.bold=bold;r.italic=italic;return
 if node.name=='a':hyperlink(p,node.get_text(),node.get('href','#'));return
 if node.name=='br':p.add_run().add_break();return
 if node.name=='code':
  r=p.add_run(node.get_text());r.font.name='Consolas';r.font.size=Pt(8.5);return
 for c in node.children:inline(p,c,bold or node.name in ['b','strong'],italic or node.name in ['i','em'])
def render(el):
 if isinstance(el,NavigableString):return
 tag=el.name
 if tag in ['h2','h3','h4']:D.add_heading(el.get_text(),{'h2':1,'h3':2,'h4':3}[tag]);return
 if tag=='p':
  p=D.add_paragraph();inline(p,el);return
 if tag in ['ul','ol']:
  for li in el.find_all('li',recursive=False):
   p=D.add_paragraph(style='List Bullet' if tag=='ul' else 'List Number');inline(p,li)
  return
 if tag=='table':
  rows=el.find_all('tr');cols=max(len(r.find_all(['td','th'],recursive=False)) for r in rows);t=D.add_table(rows=0,cols=cols);t.style='Light Shading Accent 1'
  for ri,row in enumerate(rows):
   cells=t.add_row().cells
   for ci,c in enumerate(row.find_all(['td','th'],recursive=False)):
    p=cells[ci].paragraphs[0];inline(p,c)
    for r in p.runs:r.font.size=Pt(8)
    cells[ci].vertical_alignment=0
   if ri==0:
    rep=OxmlElement('w:tblHeader');t.rows[ri]._tr.get_or_add_trPr().append(rep)
   # Keep each ordinary row together, while allowing long tables to paginate.
   cant=OxmlElement('w:cantSplit');t.rows[ri]._tr.get_or_add_trPr().append(cant)
  D.add_paragraph();return
 if tag=='pre':
  p=D.add_paragraph();r=p.add_run(el.get_text());r.font.name='Consolas';r.font.size=Pt(8);return
 if tag=='blockquote':
  p=D.add_paragraph(style='Intense Quote');inline(p,el);return
 if tag=='hr':return
 for c in el.children:render(c)
for el in soup.children:render(el)
D.add_heading('Визуальные доказательства',1)
for name,title,caption in figs:
 D.add_heading(title,2);D.add_picture(str(B/'evidence/screens'/name),height=Cm(13.0));D.add_paragraph(caption);D.add_paragraph('evidence/screens/'+name,'Caption')
D.core_properties.title='ПОТОЛКОВО — консолидированный аудит и ТЗ v6';D.core_properties.subject='Воронка калькулятора и умное освещение EKS';D.core_properties.author='Аудит проекта ПОТОЛКОВО'
D.save(B/'POTOLKOVO-FINAL-AUDIT-TZ-V6.docx')
print('HTML and DOCX ready')
