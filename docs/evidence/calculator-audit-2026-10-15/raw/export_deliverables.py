from pathlib import Path
import json,re,html,base64,hashlib,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
R=Path('/home/user/repo');O=Path(__file__).parent;C=R/'docs/current'
def clean(t):return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m:m[1]+' ('+m[2]+')',t).replace('**','').replace('`','')
def doc(file,name,title):
 d=Document();sec=d.sections[0];sec.left_margin=sec.right_margin=Inches(.72);sec.top_margin=sec.bottom_margin=Inches(.7)
 for st in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  d.styles[st].font.name='Arial';d.styles[st].font.color.rgb=RGBColor.from_string('173C36')
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7);d.styles['Title'].font.size=Pt(30)
 sec.header.paragraphs[0].text='ПОТОЛКОВО / КАЛЬКУЛЯТОР / 15.10.2026'
 foot=sec.footer.paragraphs[0];foot.text='Preview · внутренний документ · ';f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
 d.add_picture(str(R/'design/brand/v8/exports/potolkovo-horizontal-ink.png'),width=Inches(2.5));d.add_paragraph(title,'Title');d.add_paragraph('Три шага · выбор решения · комнаты · LED-зоны · корректные роли оборудования')
 d.add_paragraph('67 зафиксированных состояний браузера, исходный flow и данные каталога. Адресные исправления ролей и снятие артикула отделены от будущего изменения UX/BOM. Ставки не переписаны; конверсия реальных лидов не измерялась.');d.add_page_break()
 lines=file.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i]
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].startswith('|'):
    row=[clean(x.strip()) for x in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[:\- ]+',x or ' ') for x in row):rows.append(row)
    i+=1
   table=d.add_table(rows=0,cols=max(map(len,rows)));table.style='Light Shading Accent 1'
   for k,row in enumerate(rows):
    cells=table.add_row().cells
    for j,v in enumerate(row):cells[j].text=v
    if k==0:el=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(el)
   continue
  if line.startswith('#'):d.add_heading(clean(line.lstrip('#').strip()),min(3,len(line)-len(line.lstrip('#'))))
  elif line.startswith('- '):d.add_paragraph(clean(line[2:]),style='List Bullet')
  elif line.strip():d.add_paragraph(clean(line))
  i+=1
 d.core_properties.title=title;d.core_properties.subject='Канонический источник: docs/current/'+file.name;d.save(O/name)
doc(C/'CALCULATOR-AUDIT.md','POTOLKOVO-CALCULATOR-AUDIT.docx','Калькулятор: где теряется понятность')
doc(C/'CALCULATOR-UX-TZ.md','POTOLKOVO-CALCULATOR-UX-TZ.docx','Калькулятор: целевой UX-контракт')
rows=json.loads((O/'evidence/fixture-audit.json').read_text());w=Workbook();guide=w.active;guide.title='Читать сначала'
for row in [
 ['ПОТОЛКОВО · 15.10.2026','Аудит ролей светильников и точечные исправления'],
 ['Основание','547 записей прежнего snapshot, уточнение владельца, runtime-предикаты. Это не свежий прайс и не физическая проверка всех моделей.'],
 ['Что сделано','12 адресных ролей: 5 управления треками, 7 дополнений. 9 ложных точек исключены из счётчиков. 48 LED-панелей сохранены.'],
 ['Снятый SKU','0У-00001341 / offer2021 отсутствует в активном index/prefill/API/подборе. Неактивный архив нужен только старым обращениям.'],
 ['Что не сделано','Новый flow сценариев/комнат/LED-зон и инженерный подбор — отдельное ТЗ, не результат этой переклассификации.'],
 ['Совместимость','TRACK_CONTROL не подтверждает питание самого устройства и не делает модель универсальной для COLIBRI/CLARUS. Нужны точные pairing/паспорта.'],
 ['Цена','Числа из snapshot нужны для воспроизводимости, не являются подтверждённым предложением покупки.'],
]:guide.append(row)
guide.column_dimensions['A'].width=27;guide.column_dimensions['B'].width=113
headers=['Offer ID','Артикул','Наименование','Kind до','Kind после','Роль','Система/семейство','Активен','Точка по старому правилу','Точка после','Световая панель после','Раздел после','Требует лампу','Цена snapshot','Единица','Статус проверки']
def sheet(name,subset):
 s=w.create_sheet(name);s.append(headers);s.freeze_panes='D2';s.auto_filter.ref=f'A1:P{len(subset)+1}';s.sheet_view.showGridLines=False
 widths=[12,24,70,24,26,24,22,14,24,20,25,26,20,21,18,65]
 from openpyxl.utils import get_column_letter
 for i,width in enumerate(widths,1):s.column_dimensions[get_column_letter(i)].width=width
 for p in subset:
  s.append([p['offerId'],p['vendorCode'],p['name'],p['beforeKind'],p['afterKind'],p['role'] or '',p['system'],'Да' if p['active'] else 'Нет','Да' if p['beforePoint'] else 'Нет','Да' if p['afterPoint'] else 'Нет','Да' if p['afterPanel'] else 'Нет',p['category'] or 'не определён',p['requiredLamp'] or 'нет',p['priceRub'],p['unit'],p['review']]);s.row_dimensions[s.max_row].height=60
 s.print_title_rows='1:1'
 for cell in s[1]:cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=10);cell.fill=PatternFill('solid',fgColor='173C36');cell.alignment=Alignment(wrap_text=True)
 return s
sheet('Адресные изменения', [p for p in rows if p['role'] or p['vendorCode']=='0У-00001341'])
sheet('Источники света', [p for p in rows if p['active'] and (p['afterPoint'] or p['afterKind'] in ['TRACK_FIXTURE','CHANDELIER','LAMP'])])
sheet('Нет browse-раздела', [p for p in rows if p['active'] and p['category'] is None])
sheet('Все записи',rows)
for s in w:
 for row in s.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name='Arial',size=10,color='173C36');cell.alignment=Alignment(wrap_text=True,vertical='top');cell.fill=PatternFill('solid',fgColor='F6F5EF' if cell.row%2==0 else 'EDF0E6')
 if s==guide:
  for row in s:s.row_dimensions[row[0].row].height=55
w.save(O/'POTOLKOVO-FIXTURE-ROLE-AUDIT.xlsx')
# Compact, offline owner-facing overview. Source findings remain in Markdown/DOCX.
def uri(p,m):return 'data:'+m+';base64,'+base64.b64encode(p.read_bytes()).decode()
def picture(name,alt):
 p=O/'screenshots'/name;return '<img src="'+uri(p,'image/jpeg')+'" alt="'+html.escape(alt)+'" loading="lazy">'
a=(C/'CALCULATOR-AUDIT.md').read_text();items=[]
for m in re.finditer(r'^### (CALC-A\d+.*?)\n(.*?)(?=^### |^## |\Z)',a,re.M|re.S):
 body=html.escape(m[2].strip());body=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',body);body=body.replace('\n\n','</p><p>').replace('\n','<br>');items.append('<details><summary>'+html.escape(m[1])+'</summary><div><p>'+body+'</p></div></details>')
css='''*{box-sizing:border-box}body{margin:0;background:#f6f5ef;color:#173c36;font:15px/1.7 Arial,sans-serif}header{padding:48px max(22px,7vw);background:#173c36;color:#f6f5ef}header img{width:250px;max-width:100%}header small{display:block;color:#d7ed89;letter-spacing:1.6px;margin-top:24px}h1{font-size:clamp(36px,5.4vw,62px);line-height:1.1;font-weight:400;letter-spacing:-2px;max-width:880px}header p{max-width:800px;color:#dbe1d8}main{max-width:1200px;margin:auto;padding:30px 22px 60px}h2{font-size:32px;font-weight:500;line-height:1.2;margin:40px 0 20px}h3{font-size:21px;font-weight:550}p{margin:0 0 15px}.notice{border-left:3px solid #a9844b;background:#edf0e6;padding:20px;margin:22px 0}.grid{display:grid;grid-template-columns:1fr 1fr;gap:22px}.three{display:grid;grid-template-columns:repeat(3,1fr);gap:18px}.box{padding:24px;border:1px solid #b6c3b4;border-radius:6px}.box strong.big{font-size:38px;display:block;line-height:1.2}.shot{padding:20px;background:#edf0e6;border-radius:5px}.shot img{display:block;width:100%;height:auto;max-width:390px;margin:auto}.caption{font-size:12px;color:#53675e;margin-top:12px}.tag{display:inline-block;background:#edf0e6;padding:6px 11px;border-radius:4px;font-size:11px}.steps{display:grid;grid-template-columns:repeat(3,1fr);gap:15px}.step{border-top:3px solid #173c36;padding:22px;background:#edf0e6}details{border:1px solid #b6c3b4;border-radius:5px;margin:12px 0}summary{padding:18px;cursor:pointer;font-weight:600}details>div{padding:0 18px 18px;font-size:14px}li{margin-bottom:10px}footer{margin-top:35px;padding-top:20px;border-top:1px solid #b6c3b4;color:#53675e;font-size:12px}code{overflow-wrap:anywhere}@media(max-width:700px){.grid,.three,.steps{grid-template-columns:1fr}header{padding:32px 22px}h2{font-size:28px}}'''
body='<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ПОТОЛКОВО · аудит калькулятора</title><style>'+css+'</style><header><img src="'+uri(R/'design/brand/v8/exports/potolkovo-horizontal-reverse.svg','image/svg+xml')+'" alt="ПОТОЛКОВО"><small>15.10.2026 / ТРИ ШАГА / PREVIEW</small><h1>Меньше догадок.<br>Понятнее следующий шаг.</h1><p>Аудит сценариев, комнат, подбора света и итоговой сметы. Адресные исправления каталога выполнены отдельно от будущего изменения flow.</p></header><main><div class="three"><div class="box"><strong class="big">67</strong>состояний и переходов зафиксировано в браузере</div><div class="box"><strong class="big">18</strong>приоритетных находок по трём шагам</div><div class="box"><strong class="big">12 + 1</strong>ролей товаров исправлено + неиспользуемый артикул снят</div></div><h2>Возвращать сценарии?</h2><div class="notice"><b>Выбор результата — да. Обязательные «Стандартный / Современный / Продвинутый» — нет.</b><p>Основные варианты потолка должны быть доступны на первом этапе. Управление светом выбирается независимо: простой потолок может быть с умным светом, парящий — с обычным выключателем.</p></div><div class="steps"><div class="step"><span class="tag">01 / ПОМЕЩЕНИЯ И ПОТОЛОК</span><h3>Что и где делаем?</h3><p>Площадь, активная комната, вид потолка и необходимые функции. Понятная проверка потолочной части.</p></div><div class="step"><span class="tag">02 / СВЕТ И УПРАВЛЕНИЕ</span><h3>Что нужно каждой зоне?</h3><p>Длины из первого этапа, лента/питание/контроллер/пульт. Без ручного повторного ввода известного.</p></div><div class="step"><span class="tag">03 / СМЕТА И ЗАЯВКА</span><h3>Состав и ограничения.</h3><p>Один итог, объяснение минимальной суммы, статусы проверки и сохранённый контекст лида.</p></div></div><p class="caption">Это целевой flow из ТЗ. Новый LED-подбор, переустройство комнат и прогресса ещё не внедрены.</p><h2>Комнаты уже есть. Проблема — в маршруте.</h2><div class="grid"><div class="box"><h3>Что работает</h3><ul><li>Первая комната и быстрые площади.</li><li>Название, добавление второй комнаты и редактирование.</li><li>Сохранение данных Кухни и Спальни при правке.</li><li>Восстановление черновика по явному выбору.</li></ul></div><div class="box"><h3>Что мешает</h3><ul><li>После изменения площади снова идут вопросы о типе/карнизе/свете.</li><li>На последующих вопросах не виден устойчивый контекст комнаты.</li><li>Удаление происходит без подтверждения/Undo.</li><li>Пять реальных базовых вопросов объявлены как восемь; проверка — 7/8.</li></ul></div></div><div class="grid"><div class="shot">'+picture('02-default-390-review.jpg','Исходная проверка комнаты: 7 из 8 и неактивный футер')+'<p class="caption">До изменения flow: сводка первого этапа, «7 из 8», несколько конкурирующих действий и неактивный футер.</p></div><div class="shot">'+picture('13-floating-light-step.jpg','Исходный подбор света после парящего потолка')+'<p class="caption">Парящий потолок без треков/спотов пока не формирует задачу LED-подбора. Это ключевая связь с новым каталоговым ТЗ.</p></div></div><h2>Что уже исправлено в каталоге</h2><div class="grid"><div class="box"><h3>Управление и дополнения</h3><p><b>5 позиций управления</b> — отдельная группа «Управление» в трековых системах.</p><p><b>7 отражателей/креплений/подвесов</b> — «Крепёж и отражатели».</p><p>Они не считаются светильниками, не закрывают план точек и не вызывают фиктивную потребность ламп. Настоящие LED-панели сохранены.</p></div><div class="box"><h3>0У-00001341 снят</h3><p>Исключён из активных index/prefill, подбора, ценовых входов и публичной выдачи.</p><p>Историческая запись остаётся неактивной для старых обращений. В восстановленном расчёте — явное предупреждение и удаление по выбору, без тихой замены другим профилем.</p></div></div><div class="notice">Классификация управления не является подтверждением универсальной совместимости COLIBRI/CLARUS. Точные пары, LED-BOM и мощность требуют паспортных данных и инженерного допуска.</div><h2>Все 18 находок</h2>'+''.join(items)+'<h2>Рекомендуемый порядок</h2><ol><li>Адресные роли и снятие SKU — текущий проверяемый кодовый проход.</li><li>Контекст комнат, редактирование с возвратом, удаление, прогресс и настоящий CTA в футере.</li><li>Данные LED/паспорта и связка room → feature → zone → power circuit → единая смета.</li><li>Наблюдаемые пользовательские проходы и ручная приёмка устройств, клавиатуры и screen reader.</li></ol><footer>Исходный UX проверен на 8c14802, в локальной production-сборке 390×844 и 1440×1000. Новые исправления каталога проходят отдельный QA. Реальные заявки не отправлялись, конверсия пользователей и инженерная совместимость всех моделей не измерялись. Main/production не выпускаются автоматически.</footer></main></html>'
assert len(items)==18
body=body.replace('Новые исправления каталога проходят отдельный QA.', 'Исправления проверены на 2775d4b: 1305 unit; 288 E2E passed / 12 skipped / 0 failed / 0 flaky. Последний отправленный коммит: 569ad55.')
(O/'POTOLKOVO-CALCULATOR-AUDIT.html').write_text(body)
print('Exported two DOCX, fixture-role workbook, offline overview with 18 findings.')
