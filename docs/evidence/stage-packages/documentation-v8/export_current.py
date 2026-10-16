from pathlib import Path
import re, json, html, zipfile, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path('/home/user/repo'); C=R/'docs/current'; O=Path(__file__).parent

def plain(text):
    text=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text)
    return re.sub(r'\*\*|`','',text)

def document(source, output, subtitle):
    d=Document();sec=d.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.75)
    for name in ['Normal','Heading 1','Heading 2','Heading 3']:
        style=d.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor.from_string('173C36')
    d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.space_after=Pt(6)
    for n,size in [('Heading 1',23),('Heading 2',16),('Heading 3',12)]:d.styles[n].font.size=Pt(size)
    header=sec.header.paragraphs[0];header.text='ПОТОЛКОВО   /   ДОКУМЕНТАЦИЯ v8.1';header.runs[0].font.size=Pt(9)
    foot=sec.footer.paragraphs[0];foot.text='Только preview · Выпуск в main — по отдельному разрешению  |  '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    in_code=False;table=[]
    def flush_table():
        nonlocal table
        if not table:return
        rows=[row for row in table if not all(re.fullmatch(r':?-+:?',col.strip()) for col in row)]
        if rows:
            t=d.add_table(rows=0, cols=len(rows[0]));t.style='Light Shading Accent 1'
            for row in rows:
                cells=t.add_row().cells
                for i,v in enumerate(row[:len(cells)]):cells[i].text=plain(v.strip())
        table=[]
    for line in source.read_text().splitlines():
        if line.startswith('|'):
            table.append(line.strip('|').split('|'));continue
        flush_table()
        if line.startswith('```'):in_code=not in_code;continue
        if not line.strip():continue
        if line.startswith('# '):
            d.add_heading(plain(line[2:]),0);d.add_paragraph(subtitle)
        elif line.startswith('### '):d.add_heading(plain(line[4:]),2)
        elif line.startswith('## '):d.add_heading(plain(line[3:]),1)
        else:
            style='List Bullet' if line.startswith('- ') else None
            p=d.add_paragraph(plain(line[2:] if style else line),style)
            if in_code:
                for run in p.runs:run.font.name='Consolas';run.font.size=Pt(9)
    flush_table();d.save(output)

document(C/'BRAND-REDESIGN-TZ-V8.md',O/'POTOLKOVO-BRAND-REDESIGN-TZ-V8.docx','Остаточный объём после внедрения основных страниц · 14 октября 2026')
document(C/'TECHNICAL-TZ-V8.md',O/'POTOLKOVO-TECHNICAL-TZ-V8.docx','Технические допуски отдельно от визуальных задач · не новый live-аудит')
document(C/'PT-CONTINUATION-TZ.md',O/'POTOLKOVO-PT-CONTINUATION-TZ.docx','Действующая очередь PT-020…049 из main · продолжение в preview · 14 октября 2026')
# Compact offline plan, not a competing specification.
brand=(C/'BRAND-REDESIGN-TZ-V8.md').read_text();sections=re.split(r'^### (B8-\d+ · .+)$',brand,flags=re.M)
cards=[]
for i in range(1,len(sections),2):
    body=sections[i+1].split('\n## ')[0]
    paras=''.join('<p>'+html.escape(plain(x))+'</p>' for x in body.strip().split('\n\n'))
    cards.append('<details><summary>'+html.escape(sections[i])+'</summary>'+paras+'</details>')
css='''*{box-sizing:border-box}body{margin:0;background:#f5f3ed;color:#173c36;font:15px/1.7 Arial,sans-serif}header{background:#173c36;color:#f5f3ed;padding:48px max(22px,7vw)}header small{letter-spacing:2px;color:#d0dcc5}h1{font-size:clamp(31px,5vw,56px);font-weight:400;line-height:1.08;letter-spacing:-1.7px;max-width:900px}header p{max-width:850px;color:#d9e2d2}main{max-width:1160px;margin:auto;padding:36px 22px}h2{font-size:27px;font-weight:500;line-height:1.2}.grid{display:grid;grid-template-columns:1fr 1fr;gap:20px;margin:28px 0}.box{background:#e6ebde;padding:24px;border-top:3px solid #789263}.box p{font-size:14px}.box b{font-size:21px;font-weight:500}details{border-bottom:1px solid #b9c6b2;padding:10px 0}summary{padding:12px 3px;cursor:pointer;font-size:18px;min-height:48px}summary:focus-visible{outline:3px solid #92723a}details p{font-size:14px;white-space:pre-line;color:#405749;padding:0 8px}.note{border-left:3px solid #829668;padding:12px 20px;margin:26px 0;background:#e6ebde}footer{margin-top:36px;border-top:1px solid #b9c6b2;padding-top:20px;font-size:12px;color:#526550}code{overflow-wrap:anywhere}@media(max-width:650px){.grid{grid-template-columns:1fr}header{padding:34px 22px}summary{font-size:16px}}'''
body=f'''<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ПОТОЛКОВО · Что осталось · v8</title><style>{css}</style><header><small>ПОТОЛКОВО / ПЛАН ЗАВЕРШЕНИЯ / v8</small><h1>Основа сайта готова.<br>PT, бренд и каталог света.</h1><p>Не новый редизайн с нуля. Действующая очередь PT-020…049, брендовые задачи и новый аудит/ТЗ каталога LED — в той же preview-ветке. Раздел лент и инженерный подбор ещё не внедрены.</p></header><main><div class="grid"><section class="box"><b>Уже внедрено</b><p>Главная, калькулятор и итог, каталог, 9 услуг, компания, консультационная страница умного света, цифровые иконки и social covers, вспомогательные страницы, отзывы и новая галерея.</p></section><section class="box"><b>Только в preview</b><p><code>feat/company-rebrand-v7-preview</code>. Main возвращён к прежней версии и автоматически выпускается в production. Нового PR/merge в main без разрешения не будет.</p></section></div><h2>Продолжение PT-020 и далее</h2><p>PT-020 сохранена: восемь ключевых файлов совпали с main f3d7b69; 35 целевых тестов повторно прошли с PostgreSQL и fake delivery. Следующая кодовая задача — PT-021: доверие IP и ограничение тела запроса. PT-024 также открыта: значение аналитического события не соответствует параметру.</p><p>Действующее ТЗ PT содержит 30 карточек PT-020…049 и сценарии S01…S19. PT-ID не заменяются T8. PT-050+ из двух аудитов различаются по источнику. Основные страницы не переделываем, уже выполненные задачи не начинаем с нуля.</p><h2>Что осталось по ребрендингу и дизайну</h2>{''.join(cards)}<div class="note"><b>Что делать следующим</b><p>Технический поток: PT-021 → приёмка остатка PT-022/023 → PT-024 и далее по зависимостям. Брендовый поток: сквозной корпоративный текст и реестр обещаний → брендбук/мастера → клиентские документы и подтверждённые материалы → финальные состояния и ручная приёмка → отдельно согласованный выпуск.</p></div><h2>Технические допуски и связь с T8</h2><p>Юридическая модель, единицы/свежесть/остатки EKS, совместимость и BOM, server quote/intent, эксплуатация и security не объявляются исправленными новым внешним видом. В техническом ТЗ сохранены 36 направлений; оно дополняет действующую PT-очередь, не отменяет её. Reload автономной корзины отложен владельцем и не заменяет дизайн.</p><h2>Документация очищена</h2><p>Удалены 209 устаревших версий файлов и дублей — 12,4 МиБ. Сохранность 1854 исходных файлов/ассетов/конфигурации проверена хэшами. Новое ТЗ содержит 8 брендовых направлений, технический реестр — 36; карта происхождения — 138 строк, включая PT-020. История Git, оригиналы вложений и резервные результаты реализации сохранены.</p><footer>14.10.2026 · Основание: актуальные STATUS и UI-STATES. Последний полный QA интерфейса — код 2aca11b: 1323 unit с PostgreSQL; 304 E2E passed / 12 skipped / 0 failed / 0 flaky. Проверки материалов и прежние этапы привязаны к своим SHA в STATUS. Это не проверка production и не подтверждение прав/SMART. Канонические документы — repo/docs/current. Этот обзор — краткое представление тех же задач.</footer></main></html>'''
(O/'POTOLKOVO-PLAN-V8.html').write_text(body)
# New specification package only; historical reports are not copied into it.
entries=[]
for p in sorted(C.iterdir()):
    if p.is_file():entries.append((p,'docs/current/'+p.name))
entries += [(R/'docs/maintenance/DOCS-CLEANUP-V8.json','docs/maintenance/DOCS-CLEANUP-V8.json')]
for p in sorted((R/'docs/evidence').rglob('*')):
    if p.is_file():entries.append((p,str(p.relative_to(R))))
# Preserve repo-relative links to the new identity kit.
for p in sorted((R/'design/brand/v8').rglob('*')):
    if p.is_file(): entries.append((p,str(p.relative_to(R))))
# Client materials retain the same repo-relative layout.
for p in sorted((R/'design/client/v8').rglob('*')):
    if p.is_file(): entries.append((p,str(p.relative_to(R))))
for rel in ['scripts/build-client-materials.py','scripts/export-client-materials-data.ts','scripts/export-client-materials.mjs']:
    entries.append((R/rel,rel))
for rel in ['design/brand/v7/typefaces/OFL.txt','design/brand/v7/typefaces/provenance.json','design/brand/v7/typefaces/Inter-CyrillicLatin.ttf','public/brand/v7/inter-cyrillic-latin.woff2','public/brand/v7/social/home.jpg','scripts/build-brand-kit.py','scripts/export-brand-kit.mjs']:
    entries.append((R/rel,rel))
for name in ['POTOLKOVO-BRAND-REDESIGN-TZ-V8.docx','POTOLKOVO-TECHNICAL-TZ-V8.docx','POTOLKOVO-PT-CONTINUATION-TZ.docx','POTOLKOVO-PLAN-V8.html']:entries.append((O/name,name))
sums=''.join(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+n+'\n' for p,n in entries);(O/'SHA256SUMS.txt').write_text(sums)
with zipfile.ZipFile(O/'POTOLKOVO-CURRENT-DOCS-V8.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p,n in entries:z.write(p,n)
    z.write(O/'SHA256SUMS.txt','SHA256SUMS.txt')
with zipfile.ZipFile(O/'POTOLKOVO-CURRENT-DOCS-V8.zip') as z:
    assert z.testzip() is None
    for p,n in entries:assert z.read(n)==p.read_bytes()
print('Exported three DOCX specifications, offline plan and verified ZIP:',(O/'POTOLKOVO-CURRENT-DOCS-V8.zip').stat().st_size,'bytes')
