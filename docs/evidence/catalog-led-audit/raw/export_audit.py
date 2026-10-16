from pathlib import Path
import json,re,hashlib,base64,html,zipfile,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
R=Path('/home/user/repo');O=Path(__file__).parent;E=R/'docs/evidence/catalog-led-audit';C=R/'docs/current'
source=json.loads((E/'led-assortment.json').read_text());meta=json.loads((E/'feed-metadata.json').read_text());data=json.loads((E/'data-audit.json').read_text());labels=json.loads((E/'label-observations.json').read_text());old=json.loads((R/'data/eks-feed2-snapshot.json').read_text())['products'];byid={p['offerId']:p for p in old}
# Readable Word copies from the canonical Markdown, not a second requirements source.
def clean(t):
 return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m:m[1]+' ('+m[2]+')',t).replace('**','').replace('`','')
def export_doc(file,out,title,subtitle):
 d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
 for n in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  d.styles[n].font.name='Arial';d.styles[n].font.color.rgb=RGBColor.from_string('173C36')
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7)
 d.styles['Title'].font.size=Pt(30)
 sec.header.paragraphs[0].text='ПОТОЛКОВО / КАТАЛОГ СВЕТА / 14.10.2026'
 footer=sec.footer.paragraphs[0];footer.add_run('Внутренний документ · preview · ')
 page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
 for run in footer.runs:run.font.size=Pt(8)
 logo=R/'design/brand/v8/exports/potolkovo-horizontal-ink.png'
 d.add_picture(str(logo),width=Inches(2.5));d.add_paragraph('АУДИТ И ПЛАН РЕАЛИЗАЦИИ',style='Subtitle');d.add_paragraph(title,'Title');d.add_paragraph(subtitle)
 d.add_paragraph('Проверены код и браузерная preview-версия, XML EKS и документация. Подтверждённые находки отделены от гипотез и будущих критериев. Изменения каталога этим документом не реализованы.')
 d.add_page_break();lines=file.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i]
  if line.startswith('|'):
   block=[]
   while i<len(lines) and lines[i].startswith('|'):
    row=[clean(c.strip()) for c in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[:\- ]+',c or ' ') for c in row):block.append(row)
    i+=1
   cols=max(len(row) for row in block);t=d.add_table(rows=0,cols=cols);t.style='Light Shading Accent 1'
   for ri,row in enumerate(block):
    cells=t.add_row().cells
    for j,v in enumerate(row):
     cells[j].text=v
     for p in cells[j].paragraphs:
      p.paragraph_format.space_after=Pt(4)
      for run in p.runs:run.font.size=Pt(9)
    if ri==0:
     pr=t.rows[-1]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');pr.append(repeat)
   continue
  if line.startswith('#'):
   h=len(line)-len(line.lstrip('#'));d.add_heading(clean(line.lstrip('#').strip()),min(3,h))
  elif line.startswith('- '):d.add_paragraph(clean(line[2:]),style='List Bullet')
  elif line.strip():d.add_paragraph(clean(line))
  i+=1
 d.core_properties.title=title;d.core_properties.subject='Канонический источник: docs/current/'+file.name
 d.save(out)
export_doc(C/'CATALOG-LED-AUDIT.md',O/'POTOLKOVO-CATALOG-AUDIT.docx','Каталог: что мешает лиду','15 приоритетных находок · измеренные пути · 72 LED-позиции · подтверждения и ограничения')
export_doc(C/'CATALOG-LED-TZ.md',O/'POTOLKOVO-CATALOG-LED-TZ.docx','Лента, питание и управление','Техническое задание · подбор по метражу · совместимость · 26 приёмочных сценариев')
# XLSX: raw facts, extracted hints, and verification are separate columns.
wb=Workbook();guide=wb.active;guide.title='Начать здесь'
def setup(name,headers,widths):
 ws=wb.create_sheet(name);ws.append(headers);ws.freeze_panes='E2';ws.sheet_view.showGridLines=False
 for idx,w in enumerate(widths,1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(idx)].width=w
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='173C36');c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10);c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=42;return ws
headers=['Offer ID','Артикул EKS сейчас','Артикул в snapshot','Тип для review','Наименование поставщика','Единица сейчас в сайте','Единица для нормализации','Цена поля XML, ₽','Основание цены','Мск (поле feed)','Брн (поле feed)','U по названию/фото','Мощность в названии, Вт','Основание Вт/м','Режим по названию','Заявленный ток в названии','А/канал на фото','А суммарно на фото','Ограничение входа на фото','Статус данных','Паспорт / версия','Совместимость с пультами','Решение BOM','Карточка EKS','Фото EKS','Что запросить']
ws=setup('72 SKU',headers,[12,23,23,25,63,20,27,20,27,15,15,27,24,28,24,25,22,22,31,28,30,35,30,50,45,75])
types={'strip':'Лента','psu':'БП','controller':'Контроллер','remote':'Пульт','controller_remote_kit':'Набор пульт+контроллер'}
for p in source:
 n=p['name'][0];prev=byid.get(p['id']);typ=p['auditType'];v=p['vendorCode'][0];nm=n.lower();w=re.search(r'(\d+(?:[.,]\d+)?)\s*[wW]',n);amps=re.search(r'(\d+)\s*[aAаА](?:\b|,)',n)
 mode='CCT' if 'cct' in nm else 'RGBW' if 'rgbw' in nm else 'RGB' if 'rgb' in nm else '5-in-1: режимы проверить' if '5в1' in nm else 'Mono' if 'монохром' in nm or typ=='strip' else '—'
 voltage='12–24 V (название)' if re.search(r'12\s*[-–]\s*24',n) else '24 V (название)' if '24v' in nm else 'не установлено'
 ch=total=connector='Не установлено';status='title_only / feed';obs=labels['WB5'] if p['id']=='12386' else labels['C01RF'] if p['id']=='12379' else None
 if obs:
  ch=obs['perChannelA'];total=obs['totalA'];status='label_observed, не паспорт';voltage=obs['voltageLabel'];connector='винтовой ≤12 A; круглый ≤5 A' if p['id']=='12379' else 'Не установлено'
  if p['id']=='12379':voltage='название 12–24; фото 5–24 V: review'
 unit='sale m / reel: подтвердить' if typ=='strip' else 'kit: целое, состав подтвердить' if typ=='controller_remote_kit' else 'pcs: целое; MOQ подтвердить'
 gap='max Вт/м по режимам, ширина, рез, питаемая длина, IP/тепло, priceUnit' if typ=='strip' else 'CV/CC, выходные U/I, derating, IP/габариты/охлаждение, dimming' if typ=='psu' else 'каналы/сумма/вход, режимы, RF family, pairing, gateway, firmware/Алиса, комплектность'
 row=[p['id'],v,prev['vendorCode'] if prev else '',types[typ],n,prev['unit'] if prev else 'нет',unit,float(p['price'][0]),'Не подтверждена: м/катушка/шт',float(p['stockMoscow']),float(p['stockBarnaul']),voltage,float(w[1].replace(',','.')) if w else '', 'Не подтверждено паспортом' if typ=='strip' else 'не Вт/м',mode,amps[1] if amps else '',ch,total,connector,status,'Требуется','Требуется точная пара/семейство','needs_review',p['url'][0],p['picture'][0],gap]
 ws.append(row);r=ws.max_row;ws.cell(r,24).hyperlink=p['url'][0];ws.cell(r,25).hyperlink=p['picture'][0];ws.row_dimensions[r].height=66
al=setup('Aliases 186',['Offer ID','Старый артикул','Новый артикул','Название snapshot','LED-группа','Решение'],[13,25,25,80,18,40]);groupids={p['id'] for p in source}
for x in data['changedVendorCodesSameOffer']:al.append([x['offerId'],x['old'],x['new'],x['name'],'Да' if x['offerId'] in groupids else 'Нет','Review: упаковка, GTIN, характеристики']);al.row_dimensions[al.max_row].height=42
pairs=setup('Управление — кандидаты',['Задача','Контроллеры EKS по названию','Пульт / интерфейс','Проверка','Статус'],[28,48,43,84,28])
for row in [
 ['Mono / яркость','C01RF / 0У-00006557','RC01RF белый/чёрный','Режим по названиям согласуется; нужны RF family/pairing и входные/выходные лимиты. На фото C01RF круглый вход 5 А.','needs_review'],
 ['Mono / смартфон','WIFI TUYA C01 / 0У-00006558','Приложение; пульт — не установлен','Точное приложение, регион, интеграции, токи; не считать Wi-Fi доказательством RF-пары.','needs_review'],
 ['CCT','C02RF / 0У-00006559','RC02RF / 0У-00006553, 06554','Проверить модель фото/пульта, зоны, pairing, максимум при одновременных каналах.','needs_review'],
 ['CCT / смартфон','WIFI TUYA C02 / 0У-00006560','Приложение; пульт — не установлен','Проверить ток по каналам, приложение, offline и Алису отдельно.','needs_review'],
 ['RGB','C03RF / 0У-00006561','RC03RF / 0У-00006555, 06556','Только кандидат; нужен документ точной совместимой пары.','needs_review'],
 ['RGB / смартфон','WIFI TUYA C03 / 0У-00006562','Приложение; пульт — не установлен','Не подменять RGBW/RGB+CCT трёхканальным режимом.','needs_review'],
 ['RGBW / несколько режимов','LM051 / 0У-00006563; WB5 / 0У-00006564','Зависит от подтверждённого режима/пары','Фото WB5: 6 А/канал, 15 А всего. Не переносить пары других брендов по похожему имени.','needs_review'],
 ['Готовый набор','EKS6A*2 Mono/CCT; EKS6A*3 RGB','Пульт заявлен в названии набора','Проверить состав коробки, SKU компонентов, пределы; не добавлять пульт повторно.','needs_review'],
 ['Обычный выключатель','Подходящий БП, схема коммутации','Контроллер может не требоваться','Не навязывать диммер при простой функции вкл/выкл. Финальная схема — специалист.','проектная проверка'],
 ['Zigbee / Matter / DALI / DMX / SPI','Подтверждённые SKU в этом feed не установлены','По проекту','Не показывать как гарантированно доступный EKS-комплект; нужны модель/паспорт/экосистема.','нет утверждённого комплекта'],
]:pairs.append(row);pairs.row_dimensions[pairs.max_row].height=66
ex=setup('Примеры мощности',['Пример','LED, Вт','K','D','Требуемый номинал, Вт','Кандидат только по ваттам','Важно'],[36,16,10,10,26,32,82])
for x in json.loads((E/'engineering-examples.json').read_text()):ex.append([x['case'],x['ledLoadW'],1.25,1,x['requiredW'],x['candidateNominalW'],'Учебный расчёт; p_max подтверждён условием примера, P_aux=0. Не утверждённый EKS BOM.'])
ex.append(['Один выход WB5 по фото',144,'—','—','24 V × 6 A','не суммарные 360 W','Лимит каждого выхода и общего тока проверяются отдельно. Нужен паспорт конкретной ревизии.'])
ex.append(['C01RF: круглый вход по фото','—','—','—','≤5 A','не рекламные 12 A','У винтового входа на фото Max12A; тип входа/разъёма должен быть частью проверки.'])
for row in ex.iter_rows(min_row=2):ex.row_dimensions[row[0].row].height=55
sources=setup('Источники',['Источник','URL','Применение / ограничения'],[28,95,110])
source_rows=[['XML EKS',meta['url'],'696 offer; source date '+meta['feedDate']+' без TZ; raw SHA256 '+meta['sha256']],['Карточка WB5',labels['WB5']['source'],'Фото EKS: 6 А/канал, 15 А суммарно; label_observed, не утверждение pairing/Алисы'],['Карточка C01RF',labels['C01RF']['source'],'Фото EKS: 6/12 А, круглый вход 5 А; разные нижние границы U; 5 В не включать без паспорта'],['Arlight','https://arlight.ru/support/recommendations/','Инженерные принципы; не присвоение характеристик EKS'],['MEAN WELL LPV-100','https://www.meanwell.com/Upload/PDF/LPV-100/LPV-100-SPEC.PDF','Пример разных output/efficiency/derating. Не рекомендация LPV для объекта; ограничения паспорта сохраняются'],['MiBoxer WL5','https://miboxer.com/light/m_n/pdf/WL5/WL5_FR_V1.0.pdf','Справочная другая модель; не EKS WB5'],['BTF WB5','https://www.btf-lighting.com/products/wb5-controller','Другая торговая марка; OEM/ревизия EKS не подтверждены'],['Алиса / Matter','https://alice.yandex.ru/support/ru/smart-home/supported-matter-devices','Проверять поддерживаемые модели и транспорт; Thread не обещать по этой справке'],['Алиса / Zigbee','https://alice.yandex.ru/support/ru/smart-home/turn-on/zigbee','Точные поддерживаемые устройства/хабы, не любой Zigbee'],['EKS / статья о БП','https://eksmarket.ru/blog/kak-podobrat-blok-pitaniya-dlya-svetodiodnoy-lenty/','Не смешивать указанный в статье КПД с проектным запасом мощности']]
for row in source_rows:sources.append(row);sources.cell(sources.max_row,2).hyperlink=row[1];sources.row_dimensions[sources.max_row].height=60
for ws in list(wb)[1:]:
 ws.auto_filter.ref=ws.dimensions
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name='Arial',size=10,color='173C36');cell.alignment=Alignment(wrap_text=True,vertical='top');cell.fill=PatternFill('solid',fgColor='F6F5EF' if cell.row%2==0 else 'EDF0E6');cell.border=Border(bottom=Side(style='hair',color='B6C3B4'))
 ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
notes=[['ПОТОЛКОВО / каталог и LED','Проверка 14.10.2026 · snapshot сайта 19.04.2026, source HEAD 4239468'],['Это не утверждённая совместимость','72 позиции: 36 лент, 19 БП, 17 устройств управления/наборов. Все полные BOM остаются needs_review.'],['Откуда числа','Raw XML EKS. Цены, остатки и названия — зафиксированный снимок, не обещание покупки/отгрузки.'],['Название ≠ паспорт','Напряжение/ватты/амперы, извлечённые из названия, отделены от паспорта. Вт в имени ленты ещё не доказанный максимум Вт/м.'],['Фото EKS','WB5/C01RF имеют отдельные наблюдения по маркировке; это label_observed. Реальная партия и режимы требуют сверки.'],['Пары управления','Таблица кандидатов, не белый список. Одинаковая частота и похожее имя не доказывают pairing.'],['Aliases','186 смен кодов при совпавшем offerId, из них 58 в LED-группе. Сопоставление не автоматическое.'],['Цена','Подтвердить saleUnit/packQty/MOQ/priceBasis. Аппаратные позиции не могут оставаться дробными метрами.'],['Проверки','26 будущих приёмочных сценариев — в ТЗ, не выполненные тесты. Browser audit воспроизвёл конкретные ошибки существующего UI.'],['Следующий шаг','Запросить паспорта EKS и состав ставок владельца; затем нормализовать данные и внедрять по ТЗ.']]
for row in notes:guide.append(row)
guide.column_dimensions['A'].width=30;guide.column_dimensions['B'].width=112;guide.sheet_view.showGridLines=False
for row in guide:
 guide.row_dimensions[row[0].row].height=54
 for cell in row:cell.font=Font(name='Arial',size=11,color='173C36',bold=cell.column==1);cell.alignment=Alignment(wrap_text=True,vertical='top');cell.fill=PatternFill('solid',fgColor='EDF0E6' if cell.column==1 else 'F6F5EF')
wb.save(O/'POTOLKOVO-EKS-LED-MATRIX.xlsx')
# Source ledger copied into the repository for future implementation.
(E/'source-ledger.json').write_text(json.dumps({'checkedDate':'2026-10-14','sources':[dict(zip(['name','url','scope'],row)) for row in source_rows]},ensure_ascii=False,indent=2)+'\n')
print('Exported audit DOCX, implementation TZ DOCX and six-sheet EKS workbook.')
