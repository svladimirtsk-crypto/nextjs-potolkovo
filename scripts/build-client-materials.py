#!/usr/bin/env python3
"""B8-003/004: blank editable templates and an unverified photo register.
Design tooling only. No pricing calculation, website changes or publication.
Requires python-docx, openpyxl and Pillow. Regenerate explicitly, never over client-filled files.
"""
from pathlib import Path
from io import BytesIO
import json,html,base64,hashlib,subprocess,argparse
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from PIL import Image,ImageOps
R=Path(__file__).resolve().parent.parent;D=R/'design/client/v8';E=D/'editable'
a=argparse.ArgumentParser();a.add_argument('--refresh-templates',action='store_true');args=a.parse_args()
if E.exists() and list(E.iterdir()) and not args.refresh_templates:raise SystemExit('Refusing to replace templates: use --refresh-templates only for unfilled source templates.')
E.mkdir(exist_ok=True,parents=True)
spec=json.loads((D/'templates.json').read_text());data=json.loads((D/'source-data.json').read_text());c=data['contacts']
logo=R/'design/brand/v8/exports/potolkovo-horizontal-ink.png'
def esc(s):return html.escape(str(s),quote=True)
def uri(p,m):return 'data:'+m+';base64,'+base64.b64encode(p.read_bytes()).decode()
def block_html(b):
 t=b['type']
 if t=='fields':return '<div class="fields">'+''.join('<div class="field"><small>'+esc(k)+'</small><span>'+esc(v)+'</span></div>' for k,v in b['items'])+'</div>'
 if t=='table':return '<table><thead><tr>'+''.join('<th scope="col">'+esc(v)+'</th>' for v in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(v)+'</td>' for v in row)+'</tr>' for row in b['rows'])+'</tbody></table>'
 if t=='heading':return '<h2>'+esc(b['text'])+'</h2>'
 if t=='text':return '<p>'+esc(b['text'])+'</p>'
 if t=='notice':return '<aside class="notice">'+esc(b['text'])+'</aside>'
 if t=='checks':return '<ul class="checks">'+''.join('<li>'+esc(x)+'</li>' for x in b['items'])+'</ul>'
 if t=='amount':return '<div class="amount"><small>'+esc(b['label'])+'</small><strong>'+esc(b['value'])+'</strong><p>'+esc(b['note'])+'</p></div>'
 if t=='shots':return '<div class="shots">'+''.join('<div class="shot"><span class="number">'+esc(n)+'</span><h2>'+esc(title)+'</h2><p>'+esc(text)+'</p></div>' for n,title,text in b['items'])+'</div>'
 raise ValueError(t)
sections=[];n=0
for doc in spec['documents']:
 for local,page in enumerate(doc['pages'],1):
  n+=1
  sections.append('<section class="page" id="'+doc['id']+'-'+str(local)+'"><header class="brandhead"><img class="logo" src="'+uri(logo,'image/png')+'" alt="ПОТОЛКОВО"><div class="edition">КЛИЕНТСКИЕ МАТЕРИАЛЫ<br>v'+spec['version']+' · ШАБЛОН</div></header><div class="audience">'+esc(doc['audience'])+'</div><div class="eyebrow">'+esc(page['eyebrow'])+'</div><h1>'+esc(page['title'])+'</h1><p class="intro">'+esc(page['intro'])+'</p>'+''.join(block_html(b) for b in page['blocks'])+'<footer class="footer"><div class="line"><span>'+esc(c['phoneDisplay'])+' · '+esc(c['emailDisplay'])+'</span><span>'+f'{n:02d} / 10'+'</span></div><span class="draft">Шаблон для согласования · Не отправлять до заполнения и проверки</span></footer></section>')
font=uri(R/'public/brand/v7/inter-cyrillic-latin.woff2','font/woff2')
css='@font-face{font-family:Inter;src:url('+font+');font-weight:100 900;font-display:swap}'+(D/'materials.css').read_text()
body='<!doctype html><html lang="ru"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>ПОТОЛКОВО · клиентские материалы</title><style>'+css+'</style><body><nav class="screen-nav" aria-label="О комплекте"><h2>Клиентские материалы · v8.2</h2><p>Пять редактируемых шаблонов и внутренний реестр проектов. Не договоры и не подтверждённое портфолио.</p></nav>'+''.join(sections)+'</body></html>'
(D/'POTOLKOVO-CLIENT-MATERIALS.html').write_text(body)
# Editable Word files use the same content and logo; system Arial for portable editing.
def shade(cell,color='EDF0E6'):
 pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);pr.append(sh)
def table(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 available=6.82
 for col in t.columns:col.width=Inches(available/len(headers))
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x;shade(t.rows[0].cells[i])
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=x
 for row in t.rows:
  trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(4)
    for run in p.runs:run.font.size=Pt(9.5)
 return t
for ds in spec['documents']:
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
 sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.72)
 sec.header_distance=sec.footer_distance=Inches(.28)
 for style in ['Normal','Title','Heading 1','Heading 2']:
  d.styles[style].font.name='Arial';d.styles[style].font.color.rgb=RGBColor.from_string('173C36')
 d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.12
 d.styles['Title'].font.size=Pt(26);d.styles['Title'].paragraph_format.space_after=Pt(11)
 d.styles['Heading 2'].font.size=Pt(13)
 hp=sec.header.paragraphs[0];hp.add_run('ПОТОЛКОВО / '+ds['name'].upper()).font.size=Pt(8)
 fp=sec.footer.paragraphs[0];fp.add_run(c['phoneDisplay']+'  ·  '+c['emailDisplay']+'  |  ').font.size=Pt(8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
 fp.add_run('\nШаблон для согласования. Не отправлять до заполнения и проверки.').font.size=Pt(8)
 for idx,page in enumerate(ds['pages']):
  if idx:d.add_page_break()
  if idx==0:d.add_picture(str(logo),width=Inches(2.0))
  p=d.add_paragraph(page['eyebrow']+' · '+ds['audience']);p.paragraph_format.space_after=Pt(8)
  for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('53675E')
  d.add_paragraph(page['title'].replace('\n',' '),'Title');d.add_paragraph(page['intro'])
  for b in page['blocks']:
   t=b['type']
   if t=='fields':table(d,['Поле','Заполнить'],b['items'])
   elif t=='table':table(d,b['headers'],b['rows'])
   elif t=='heading':d.add_heading(b['text'],2)
   elif t=='text':d.add_paragraph(b['text'])
   elif t=='notice':
    p=d.add_paragraph(b['text']);p.paragraph_format.space_before=Pt(7)
    for run in p.runs:run.font.size=Pt(9);run.italic=True
   elif t=='checks':
    for item in b['items']:d.add_paragraph('☐ '+item)
   elif t=='amount':
    d.add_heading(b['label'],2);p=d.add_paragraph(b['value']);p.runs[0].bold=True;p.runs[0].font.size=Pt(17);d.add_paragraph(b['note'])
   elif t=='shots':table(d,['Кадр','Что показать'],[[n+' / '+title,text] for n,title,text in b['items']])
 d.core_properties.title='ПОТОЛКОВО · '+ds['name'];d.core_properties.subject='Рабочий шаблон, v8.2; не юридически утверждённый документ';d.save(E/(ds['id']+'.docx'))
# Actual inventory, no inferred ownership or role from file presence.
records=[]
for project in data['projects']:
 for idx,src in enumerate(project['photos'],1):
  p=R/'public'/src.lstrip('/');raw=p.read_bytes()
  with Image.open(p) as im:w,h=im.size
  records.append({'photoId':project['assetKey']+'-'+str(idx),'projectSlug':project['slug'],'path':src,'width':w,'height':h,'bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest(),'source':'existing repository photo; not ownership evidence','rightsStatus':'Не подтверждено','publicationDecision':'Не публиковать'})
assert len(records)==18 and len({x['path'] for x in records})==18
(D/'photo-inventory.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')
wb=Workbook();guide=wb.active;guide.title='Начать здесь';pr=wb.create_sheet('Проекты');ph=wb.create_sheet('Фотографии')
def sheet(ws,headers,widths):
 ws.append(['ПОТОЛКОВО / РЕЕСТР ПОДТВЕРЖДЕНИЙ']);ws.append(['Внутренний документ. Все текущие материалы не подтверждены для нового кейса.']);ws.append(headers)
 for i,width in enumerate(widths,1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=width
 ws.freeze_panes='E4';ws.auto_filter.ref=f'A3:{__import__("openpyxl").utils.get_column_letter(len(headers))}3'
 ws.sheet_view.showGridLines=False
 for cell in ws[3]:cell.fill=PatternFill('solid',fgColor='173C36');cell.font=Font(name='Arial',color='FFFFFF',bold=True,size=10);cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[3].height=40;ws.print_title_rows='1:3';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
headers=['Код проекта','Рабочее название из сайта','Роль / атрибуция','Период и источник','Площадь и источник / не публиковать','Задача и ограничения','Фактически выполнено','Тип бюджета','Сумма из источника','Источник / дата / версия бюджета','Включено','Не включено','Источники фактов','Решение владельца','Проверил','Дата проверки','Открытые вопросы']
sheet(pr,headers,[32,42,30,28,28,38,38,25,22,36,32,32,36,30,24,18,42])
for p in data['projects']:pr.append([p['slug'],p['title'],'','','','','','Не выбран','','','','','','Не согласовано','','','Подтвердить факты, роль, период и права на все фотографии.']);pr.row_dimensions[pr.max_row].height=68
headers=['Миниатюра (не оригинал)','ID фото','Код проекта','Путь в репозитории','Автор','Источник оригинала / доказательство','Право публикации','Разрешённые каналы','Идентифицирующие детали','Основания / согласия / обработка','Фактическая подпись','Решение','Проверил','Дата','Ширина px','Высота px','SHA256','Байт']
sheet(ph,headers,[22,18,32,31,25,42,25,30,32,40,40,28,25,18,14,14,68,16])
thumbs=[]
for rec in records:
 ph.append(['',rec['photoId'],rec['projectSlug'],rec['path'],'','','Не подтверждено','','Не проверено','','','Не публиковать','','',rec['width'],rec['height'],rec['sha256'],rec['bytes']]);row=ph.max_row;ph.row_dimensions[row].height=78
 with Image.open(R/'public'/rec['path'].lstrip('/')) as im:
  image=ImageOps.exif_transpose(im).convert('RGB');image.thumbnail((132,92));buf=BytesIO();image.save(buf,format='PNG');buf.seek(0)
  thumb=XLImage(buf);thumbs.append(buf);ph.add_image(thumb,f'A{row}')
for col in ['O','P','Q','R']:ph.column_dimensions[col].hidden=True
for ws in [pr,ph]:
 ws.auto_filter.ref=f'A3:{__import__("openpyxl").utils.get_column_letter(ws.max_column)}{ws.max_row}'
 for row in ws.iter_rows(min_row=4):
  for cell in row:
   cell.font=Font(name='Arial',size=10,color='173C36');cell.alignment=Alignment(wrap_text=True,vertical='top');cell.fill=PatternFill('solid',fgColor='F6F5EF' if cell.row%2==0 else 'EDF0E6');cell.border=Border(bottom=Side(style='hair',color='B6C3B4'))
 ws.print_options.horizontalCentered=True

def choices(ws,col,values):
 dv=DataValidation(type='list',formula1='"'+','.join(values)+'"');dv.errorTitle='Выберите статус';dv.error='Используйте значение из списка. Подтверждение должно иметь основание.';dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(f'{col}4:{col}{ws.max_row}')
choices(pr,'H',['Не выбран','Фактический','Расчётный ориентир','Не публиковать'])
choices(pr,'N',['Не согласовано','На проверке','Согласовано владельцем'])
choices(ph,'G',['Не подтверждено','Подтверждено','Запрещено'])
choices(ph,'I',['Не проверено','Нет идентифицирующих деталей','Нужна проверка'])
choices(ph,'L',['Не публиковать','На проверке','Согласовано владельцем'])
guide.column_dimensions['A'].width=33;guide.column_dimensions['B'].width=105;guide.sheet_view.showGridLines=False
notes=[('Назначение','6 текущих примеров и 18 фото. Это инвентаризация файлов, не подтверждённое портфолио.'),('Источник','content/home-assets.ts и homepage.proof.items. Рабочие названия переписаны из кода, не подтверждены владельцем.'),('Заполнять','Поля автора, прав, роли, периода, выполненных работ и источников. Исходные документы хранить в защищённой папке, не в публичном Git.'),('Бюджет','Только из подтверждённого источника. Здесь нет ценовых формул и автоматического расчёта стоимости.'),('Статусы','Выбор «Подтверждено» не заменяет основание. Решение о публикации принимает владелец после проверки всех фактов и фотографий.'),('Блокеры','Не выпускать новый кейс без источника фото, прав/каналов, проверки идентифицирующих деталей, роли, периода, задачи, работ и решения владельца.'),('Площадь и цена','Если нет достоверных сведений, поле не публикуется. Не подменять исторический счёт сегодняшним калькулятором.'),('Миниатюры','Только для поиска кадра. Оригиналы не изменены. Ширина/высота/хэш находятся в скрытых столбцах O–R.'),('Обновление','Не перезапускать генератор поверх заполненного реестра. Сохранить заполненную копию отдельно; source templates обновлять явной командой.'),('Публикация','Этот файл и комплект не создают case-pages, не меняют галерею и не выдают автоматический допуск.'),('Конфиденциальность','Не переносить в публичные поля точный адрес, документы клиента, телефоны или другую избыточную информацию.'),('Правила','Смотрите 04-PASSPORT-PROEKTA.docx и 05-FOTOBRIF.docx в editable/.')]
for k,v in notes:guide.append([k,v])
for row in guide:
 for cell in row:cell.font=Font(name='Arial',size=11,color='173C36');cell.alignment=Alignment(wrap_text=True,vertical='top');cell.fill=PatternFill('solid',fgColor='EDF0E6' if cell.column==1 else 'F6F5EF')
 guide.row_dimensions[row[0].row].height=52
wb.save(E/'06-REESTR-PROEKTOV-I-FOTO.xlsx')
manifest={'version':spec['version'],'sourceProductCommit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip(),'contactsSource':'content/contacts.ts','legalFieldsCopied':False,'pricingCalculation':False,'automaticSiteExport':False,'documents':5,'pdfExpectedPages':10,'projects':6,'photos':18,'publication':'none; all rights and project facts pending','sourceHashes':{p:hashlib.sha256((R/p).read_bytes()).hexdigest() for p in ['content/contacts.ts','content/home-assets.ts','content/homepage.ts','content/pricing.ts','lib/calculator/types.ts']}}
(D/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print('5 editable DOCX, 1 XLSX, 10 HTML pages, 18 photo fingerprints. No pricing engine or publication.')
